"""Deterministic synthetic data. Brightwater Facilities Group is fictional, as is every
person, site, address and number below.

The workbook is seeded with one deliberate property, modelled on a pattern that shows up
in real operations workbooks: the stored Action Plan text does NOT follow the formula its
owner documents for it. `reconcile.py` holds the documented version; this file holds what
the workbook "actually does". The gap between them is why G4 exists.
"""
from __future__ import annotations

import datetime as dt
import json
import random
from pathlib import Path

from openpyxl import Workbook as XlsxWorkbook

from . import fmt
from .workbook import (COL, LEAD_COL, LEAD_HEADER_ROW, LEAD_FIRST_ROW, SHEET_LEADS, SHEET_SITES,
                       SITE_FIRST_ROW, SITE_HEADER_ROWS, WEEK_CELL)

WEEK_ENDING = dt.date(2026, 9, 13)
DOMAIN = "brightwater.example"
QA_TARGET = 0.10          # QA inspection hours as a share of service hours
QA_HIGH = 0.20
TRAIN_TARGET = 2.0        # tenant-training hours per site per month
REMOTE_MAX = 0.25

LAST = ["Abara", "Brandt", "Calloway", "Dimitrov", "Eklund", "Farrow", "Galvez", "Haldane",
        "Ibarra", "Jansson", "Kovac", "Lindqvist", "Marchetti", "Nakamura", "Okafor", "Pruitt",
        "Quintero", "Rasmussen", "Salgado", "Thackeray", "Ueda", "Varga", "Whitlock", "Yilmaz"]
FIRST = ["Imani", "Tobias", "Renata", "Pavel", "Signe", "Desmond", "Lucia", "Fergus",
         "Marisol", "Anders", "Mila", "Henrik", "Carla", "Kenji", "Adaeze", "Wendell",
         "Paloma", "Soren", "Beatriz", "Rupert", "Aiko", "Laszlo", "Odette", "Emre"]
SITE_A = ["Alder", "Birch", "Cedar", "Dunmore", "Eastgate", "Fenwick", "Garrison", "Harlow",
          "Ironbridge", "Juniper", "Kestrel", "Larkspur", "Meridian", "Northfield", "Oakhurst",
          "Pinecrest", "Quarry", "Redfern", "Stonebrook", "Tamarack", "Underhill", "Vantage",
          "Westmark", "Yarrow", "Ashby", "Bellweather", "Copperfield", "Dovetail", "Elmstead",
          "Foxhollow"]
SITE_B = ["Plaza", "Commons", "Tower", "Exchange", "Annex", "Court", "Works", "Landing", "Square"]
REGIONS = ["NE", "SE", "MW", "SW", "NW"]

# Leads deliberately missing from the workbook's Lead View, so resolution has to fall
# through to the directory - where one resolves, one is ambiguous, one is absent.
VIA_DIRECTORY = "Okafor Adaeze"
AMBIGUOUS = "Salgado Beatriz"
ABSENT = "Varga Laszlo"


def stored_action_plan(row: dict, month: str) -> str:
    """What the workbook's Action Plan column contains. Note the training line: a flat
    2-hour target, no authorized-hours cap, no round-up - and 'w/o authorization' sorts
    AFTER the remote-QA line."""
    lines = []
    s = row["qa_pct_mtd"]
    if s is not None and s > QA_HIGH:
        lines.append("Discuss high QA % of Service MTD with RM")
    if s is not None and s < QA_TARGET:
        lines.append(f"Add at least {fmt.hours(row['qa_add'])} additional QA hours to {month}")
    flat = round(TRAIN_TARGET - row["train_mtd"], 2)
    if row["train_auth_mtd"] > 0 and flat > 0:
        lines.append(f"Add {fmt.hours(flat)} Tenant training hours to {month}")
    if row["remote_pct_mtd"] is not None and row["remote_pct_mtd"] >= REMOTE_MAX:
        lines.append("Discuss excessive remote QA with RM")
    if row["train_auth_mtd"] <= 0 and row["train_mtd"] > 0:
        lines.append("Tenant training delivered w/o authorization; consider requesting authorization")
    return "\n".join("(x) " + line for line in lines)


def _site_row(rng: random.Random, site_id: str, site: str, region: str, lead: str, month: str) -> dict:
    idle = rng.random() < 0.06
    service_wk = 0.0 if idle or rng.random() < 0.08 else round(rng.uniform(4, 40), 2)
    service_mtd = 0.0 if idle else round(service_wk + rng.uniform(5, 60), 2)
    share = rng.choice([0.0, 0.03, 0.06, 0.09, 0.11, 0.13, 0.16, 0.24])
    qa_wk = round(service_wk * share * rng.uniform(0.5, 1.5), 2)
    qa_mtd = round(max(qa_wk, service_mtd * share * rng.uniform(0.7, 1.2)), 2)
    train_wk = rng.choice([0, 0, 0, 0, 0.5, 0.75, 1.0, 1.25])
    train_mtd = round(train_wk + rng.choice([0, 0, 0.25, 0.5, 1.0, 1.5, 2.0]), 2)
    auth = rng.choice([0, 0, 1.0, 1.5, 2.0, 2.0, 4.0, 4.0])
    row = {
        "site_id": site_id, "site": site, "region": region, "lead": lead,
        "service_wk": service_wk, "service_mtd": service_mtd,
        "qa_wk": qa_wk, "qa_mtd": qa_mtd,
        "qa_pct_wk": round(qa_wk / service_wk, 4) if service_wk else None,
        "qa_pct_mtd": round(qa_mtd / service_mtd, 4) if service_mtd else None,
        "qa_add": round(max(0.0, QA_TARGET * service_mtd - qa_mtd), 2),
        "train_wk": float(train_wk), "train_mtd": train_mtd,
        # The "hours to add" COLUMN honours the authorized cap. The action TEXT doesn't.
        "train_add": round(max(0.0, min(TRAIN_TARGET, auth) - train_mtd), 2) if auth > 0 else 0.0,
        "train_auth_mtd": float(auth),
        "remote_pct_mtd": round(rng.choice([0, 0.05, 0.1, 0.18, 0.27, 0.4]), 4) if qa_mtd else None,
    }
    row["action_plan"] = stored_action_plan(row, month)
    return row


def build(seed: int = 20260913) -> dict:
    rng = random.Random(seed)
    leads = [f"{last} {first}" for last, first in zip(LAST, FIRST)]
    names = [f"{a} {b}" for a in SITE_A for b in SITE_B]
    rng.shuffle(names)
    month = WEEK_ENDING.strftime("%B")

    rows, n = [], 0
    for lead in leads:
        region = rng.choice(REGIONS)
        for _ in range(rng.randint(4, 13)):
            rows.append(_site_row(rng, f"S-{10400 + n}", names[n], region, lead, month))
            n += 1

    off_sheet = {VIA_DIRECTORY, AMBIGUOUS, ABSENT}
    lead_view, roster, directory = [], {}, []
    for i, lead in enumerate(leads):
        last, first = lead.split(" ", 1)
        email = f"{first}.{last}@{DOMAIN}".lower()
        if lead not in off_sheet:
            emp = f"E{7000 + i}"
            lead_view.append((emp, lead))
            roster[emp] = email
        if lead == ABSENT:
            continue
        directory.append({"name": f"{first} {last}", "email": email, "title": "Site Lead"})
        if lead == AMBIGUOUS:
            directory.append({"name": f"{first} {last}", "email": f"{first[0]}{last}2@{DOMAIN}".lower(),
                              "title": "Accounts Payable Specialist"})
    return {"rows": rows, "lead_view": lead_view, "roster": roster, "directory": directory}


def write(data_dir: Path, seed: int = 20260913, week: dt.date = WEEK_ENDING) -> Path:
    data_dir = Path(data_dir)
    data_dir.mkdir(parents=True, exist_ok=True)
    data = build(seed)

    wb = XlsxWorkbook()
    ws = wb.active
    ws.title = SHEET_SITES
    ws["B1"] = "Site KPI Analysis (Ops Version)"
    ws["B3"] = "Week ending"
    ws[WEEK_CELL] = week
    labels = {
        "site_id": ("Site", "External ID"), "site": ("Site", "Name"), "region": ("Site", "Region Code"),
        "lead": ("Site", "Site Lead"),
        "service_wk": ("Service", "Rendered Hours"), "service_mtd": ("Service", "Rendered Hours MTD"),
        "qa_wk": ("QA", "Rendered Hours"), "qa_mtd": ("QA", "Rendered Hours MTD"),
        "qa_pct_wk": ("QA", "% of Service"), "qa_pct_mtd": ("QA", "% of Service MTD"),
        "qa_add": ("QA", "Hours to Add"),
        "train_wk": ("Tenant Training", "Rendered"), "train_mtd": ("Tenant Training", "Rendered MTD"),
        "train_add": ("Tenant Training", "Hours to Add"),
        "train_auth_mtd": ("Tenant Training", "Authorized MTD"),
        "remote_pct_mtd": ("Remote", "% of QA Rendered MTD"),
        "action_plan": ("Plan", "Action Plan"),
    }
    for key, (group, name) in labels.items():
        ws.cell(row=SITE_HEADER_ROWS[0], column=COL[key], value=group)
        ws.cell(row=SITE_HEADER_ROWS[1], column=COL[key], value=name)
    for r, row in enumerate(data["rows"], start=SITE_FIRST_ROW):
        for key, col in COL.items():
            ws.cell(row=r, column=col, value=row[key])

    lv = wb.create_sheet(SHEET_LEADS)
    lv.cell(row=LEAD_HEADER_ROW, column=LEAD_COL["employee_id"], value="Employee Id")
    lv.cell(row=LEAD_HEADER_ROW, column=LEAD_COL["employee_name"], value="Employee Name")
    for r, (emp, name) in enumerate(data["lead_view"], start=LEAD_FIRST_ROW):
        lv.cell(row=r, column=LEAD_COL["employee_id"], value=emp)
        lv.cell(row=r, column=LEAD_COL["employee_name"], value=name)

    path = data_dir / f"Site KPI Analysis (Ops Version) - {week:%Y.%m.%d}.xlsx"
    wb.save(path)

    (data_dir / "announcement.json").write_text(json.dumps({
        "from": f"ops-reporting@{DOMAIN}",
        "subject": f"Site KPI Analysis - {week.month}.{week.day:02d}.{week:%y}",
        "received": f"{week + dt.timedelta(days=1)}T16:42:00",
    }, indent=2) + "\n")
    (data_dir / "roster.json").write_text(json.dumps(data["roster"], indent=2) + "\n")
    (data_dir / "directory.json").write_text(json.dumps(data["directory"], indent=2) + "\n")
    return path
